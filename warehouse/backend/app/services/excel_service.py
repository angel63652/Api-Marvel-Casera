"""Excel (.xlsx) generation helper using openpyxl."""
import io
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _cell(value):
    """Coerce values openpyxl can't write natively to safe types."""
    if isinstance(value, (datetime, date, int, float, str)) or value is None:
        return value
    if hasattr(value, "value"):  # Enum
        return value.value
    return str(value)


def build_xlsx(sheet_name: str, headers: list[str], rows: list[list]) -> bytes:
    """Build a single-sheet workbook with a bold header row; returns bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Hoja1"

    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for row in rows:
        ws.append([_cell(v) for v in row])

    # Reasonable column widths from content length.
    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            if idx - 1 < len(row):
                max_len = max(max_len, len(str(row[idx - 1]) if row[idx - 1] is not None else ""))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_headers(filename: str) -> dict:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}
